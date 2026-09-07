# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY = "1E2761"
NAVY_DARK = "24306E"
NAVY_DARK2 = "2B3A80"
ICE = "CADCFC"
WHITE = "FFFFFF"
AMBER = "F2A900"
MUTED = "5A6B8C"
LIGHTBG = "F6F8FC"
BORDER = "E3E8F5"
ICE_TXT = "9FB2E0"
ROWALT = "EEF2FA"

EA_FONT = "맑은 고딕"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def set_ea(run, ea_name=EA_FONT):
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', ea_name)


def add_slide(bg=WHITE):
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(bg)
    return slide


def add_text(slide, x, y, w, h, text, size, color=NAVY, bold=False, italic=False,
             font="Calibri", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, url=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = rgb(color)
    set_ea(run)
    if url:
        run.hyperlink.address = url
    return box


def add_rect(slide, x, y, w, h, fill, rounded=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_oval(slide, x, y, d, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_badge(slide, x, y, d, num, fill, textcolor, size=None):
    shp = add_oval(slide, x, y, d, fill)
    tf = shp.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(num)
    f = run.font
    f.size = Pt(size or (18 if d > 0.7 else 13))
    f.bold = True
    f.name = "Calibri"
    f.color.rgb = rgb(textcolor)
    set_ea(run)
    return shp


def add_line(slide, x, y, w, color, weight_pt=3):
    h_in = weight_pt / 72.0
    add_rect(slide, x, y - h_in / 2, w, h_in, color, rounded=False)


def header(slide, tag, title, subtitle=None, dark=False):
    tagc = AMBER
    titlec = WHITE if dark else NAVY
    subc = ICE_TXT if dark else MUTED
    add_text(slide, 0.6, 0.4, 8.0, 0.4, tag, 14, tagc, bold=True)
    add_text(slide, 0.6, 0.75, 11.5, 0.65, title, 27, titlec, bold=True, font="Cambria")
    if subtitle:
        add_text(slide, 0.6, 1.38, 11.5, 0.4, subtitle, 13.5, subc)


LINK_COLOR = "1155CC"


def set_cell(cell, text, size=11.5, bold=False, color=NAVY, align=PP_ALIGN.LEFT,
             fill=WHITE, font="Calibri", url=None):
    cell.margin_left = Inches(0.12)
    cell.margin_right = Inches(0.12)
    cell.margin_top = Inches(0.06)
    cell.margin_bottom = Inches(0.06)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.name = font
    if url:
        f.color.rgb = rgb(LINK_COLOR)
        f.underline = True
    else:
        f.color.rgb = rgb(color)
    set_ea(run)
    if url:
        run.hyperlink.address = url


def add_table(slide, x, y, w, col_widths, row_heights, header_row, rows, header_fill=NAVY,
              header_color=WHITE, body_size=11.5):
    n_rows = len(rows) + 1
    n_cols = len(col_widths)
    total_h = sum(row_heights)
    gshape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(total_h))
    table = gshape.table
    table.first_row = False
    table.horz_banding = False
    for i, cw in enumerate(col_widths):
        table.columns[i].width = Inches(cw)
    for i, rh in enumerate(row_heights):
        table.rows[i].height = Inches(rh)
    for j, htext in enumerate(header_row):
        set_cell(table.cell(0, j), htext, size=12.5, bold=True, color=header_color,
                 fill=header_fill, align=PP_ALIGN.LEFT if j > 0 else PP_ALIGN.LEFT)
    for i, row in enumerate(rows):
        fill = WHITE if i % 2 == 0 else ROWALT
        for j, cell_data in enumerate(row):
            if isinstance(cell_data, tuple):
                text, url = cell_data
            else:
                text, url = cell_data, None
            bold = (j == 0)
            color = NAVY if j == 0 else MUTED
            set_cell(table.cell(i + 1, j), text, size=body_size, bold=bold, color=color,
                     fill=fill, url=url)
    return table


def field_slide(tag, title, subtitle, rows):
    s = add_slide(LIGHTBG)
    header(s, tag, title, subtitle)
    labels = ["자격증", "필요 역량", "취업 시 필요한 것", "대표 기업", "신입 초봉 수준", "선호 학력"]
    table_rows = list(zip(labels, rows))
    add_table(
        s, 0.6, 1.9, 12.13, [2.15, 9.98],
        [0.78] * len(table_rows), ["구분", "내용"], table_rows,
        header_fill=NAVY, header_color=WHITE, body_size=12
    )
    return s


# =================================================================
# Slide 1: Title
# =================================================================
s = add_slide(NAVY)
add_oval(s, 10.6, 4.7, 4.2, NAVY_DARK)
add_oval(s, 11.6, 5.9, 2.6, NAVY_DARK2)
add_oval(s, -1.2, -1.4, 3.6, NAVY_DARK)
add_text(s, 0.9, 2.1, 10.5, 1.0, "군 복무 ~ 졸업까지", 40, WHITE, bold=True, font="Cambria")
add_text(s, 0.9, 2.95, 10.5, 0.7, "컴퓨터공학과 진로 로드맵 (상세판)", 24, ICE, bold=True, font="Cambria")
add_text(s, 0.9, 3.7, 10.0, 0.5, "인하대학교 컴퓨터공학과  ·  2026.09 잔여 복무 ~ 졸업", 15, ICE_TXT)
add_text(s, 0.9, 4.15, 10.0, 0.5,
         "학년별 로드맵 + 인턴십 프로그램 + 분야별 자격증·역량·기업·초봉·선호학력",
         13, ICE_TXT)
add_text(s, 0.9, 6.55, 9.0, 0.4, "채용 시장 데이터 기반 정리 (2025~2026 조사 기준, 추정치 포함)", 12, "8494C4", italic=True)

# =================================================================
# Slide 2: Overview timeline
# =================================================================
s = add_slide(WHITE)
add_text(s, 0.6, 0.45, 10.0, 0.6, "전체 로드맵 한눈에 보기", 28, NAVY, bold=True, font="Cambria")

stages = [
    ("잔여 복무", "~ 2026.11"),
    ("전역 공백기", "여행 · 알바"),
    ("2학년", "전공 기초"),
    ("3학년", "심화 · 인턴"),
    ("4학년 1학기", "포트폴리오"),
    ("4학년 2학기", "본격 취업"),
]
n = len(stages)
start_x, end_x, line_y = 0.9, 12.4, 2.95
gap = (end_x - start_x) / (n - 1)
add_line(s, start_x, line_y, end_x - start_x, ICE, 4)

d = 0.6
for i, (t, sub) in enumerate(stages):
    cx = start_x + gap * i
    fill = AMBER if i == n - 1 else NAVY
    add_badge(s, cx - d / 2, line_y - d / 2, d, i + 1, fill, WHITE, size=14)
    label_w = 2.1
    add_text(s, cx - label_w / 2, line_y + 0.5, label_w, 0.4, t, 13, NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, cx - label_w / 2, line_y + 0.9, label_w, 0.4, sub, 10.5, MUTED, align=PP_ALIGN.CENTER)

add_text(s, 0.9, 4.05, 11.5, 0.4,
         "이 로드맵에서 추가로 다루는 내용", 15, NAVY, bold=True)
extra_items = [
    "학년/시기별 준비 로드맵 (Stage 1~6)",
    "국내 대표 인턴십 · 채용연계 교육 프로그램 7종",
    "인턴 · 채용 정보 사이트 8곳 (바로가기 링크)",
    "분야별 상세: 백엔드 / AI·데이터 / 정보보안 / 클라우드·DevOps",
    "대기업 신입 초봉 스냅샷 + 종합 우선순위",
]
top = 4.55
for i, it in enumerate(extra_items):
    y = top + i * 0.42
    add_badge(s, 0.9, y, 0.28, "•", NAVY, WHITE, size=10)
    add_text(s, 1.35, y - 0.03, 10.8, 0.4, it, 13, MUTED)

# =================================================================
# Content stage slides
# =================================================================
def content_slide(tag, title, period, rows, footer=None):
    s = add_slide(LIGHTBG)
    header(s, tag, title, period)
    top = 2.05
    bottom = 6.55
    row_h = (bottom - top) / len(rows)
    for i, (h, b) in enumerate(rows):
        y = top + row_h * i
        add_badge(s, 0.6, y + (row_h - 0.55) / 2, 0.55, i + 1, NAVY, WHITE, size=15)
        add_text(s, 1.45, y + 0.02, 10.8, 0.35, h, 16, NAVY, bold=True)
        add_text(s, 1.45, y + 0.4, 11.0, row_h - 0.45, b, 12.5, MUTED)
    if footer:
        add_rect(s, 0.6, 6.75, 12.1, 0.02, BORDER)
        add_text(s, 0.6, 6.85, 12.1, 0.5, footer, 12, MUTED, italic=True)
    return s


content_slide(
    "STAGE 01", "잔여 복무기간", "지금 ~ 2026년 11월 전역",
    [
        ("CS 기본기 복습", "자료구조 · 운영체제 · 네트워크 등 책으로 개념 정리 (실습 없이도 가능)"),
        ("코딩 감각 유지", "휴대폰/PC 사용 가능 시 백준·프로그래머스로 간단한 문제 풀이"),
        ("정보처리기사 필기 (선택)", "일정이 맞으면 필기만 시도 — 실기는 전역 후 환경에서 준비"),
    ],
    footer="목표는 '새로운 공부'가 아니라 '감을 잃지 않는 것'"
)

content_slide(
    "STAGE 02", "전역 ~ 복학 공백기", "여행 · 아르바이트 계획 중 (약 2~3개월)",
    [
        ("자격증 · 스펙은 잠시 보류", "정보처리기사 등 시험 일정이 이 기간과 맞지 않음 — 다음 학기 방학으로 미루기"),
        ("하루 30분~1시간, 코딩 감각만 유지", "여행 중에도 가능한 수준으로 가볍게. 목표는 실력 향상이 아니라 복학 워밍업"),
        ("여행 · 알바에 집중", "사회 복귀 리프레시 시간으로 활용. 무리한 계획보다 충분한 재충전이 우선"),
    ],
    footer="이 시기를 억지로 채우려 하지 않는 것이 오히려 전략"
)

content_slide(
    "STAGE 03", "2학년 — 전공 기초 + 코딩 습관", "가장 중요한 기반을 만드는 시기",
    [
        ("전공 기초 탄탄히", "자료구조 · 알고리즘 · 이산수학 · 객체지향프로그래밍 — 개념을 코드로 짤 수 있는지가 핵심"),
        ("코딩테스트 입문", "백준 · 프로그래머스로 매주 꾸준히 (Bronze → Silver). 3~4학년에 몰아서 하면 늦음"),
        ("동아리 · 깃허브 습관", "알고리즘 스터디 등 동료 만들기 + 작은 프로젝트라도 깃허브에 커밋하는 습관"),
    ]
)

content_slide(
    "STAGE 04", "3학년 — 심화 + 분야 결정 + 인턴 도전", "가장 바쁘고 가장 중요한 학년 (인턴십 목록 9p 참고)",
    [
        ("전공 심화 + 분야 1개 특화", "운영체제 · 네트워크 · DB · 소프트웨어공학 + 백엔드/AI데이터/보안/클라우드 중 방향 설정"),
        ("자격증 1~2개", "정보처리기사, SQLD, AWS 자격증 중 선택 — 분야별 상세는 11~14p 참고"),
        ("방학 인턴십 적극 지원", "3학년 인턴 경험이 4학년 취업 성패를 가르는 경우가 많음. 광탈해도 계속 지원"),
        ("포트폴리오 프로젝트 착수", "배포까지 완료한 프로젝트 1개 — 자격증보다 채용에 직접적 영향"),
    ]
)

content_slide(
    "STAGE 05", "4학년 1학기 — 포트폴리오 완성", "채용연계형 인턴 총력",
    [
        ("캡스톤 디자인 = 포트폴리오 핵심", "학점용이 아니라 이력서에 바로 쓸 수 있는 완성도로 마무리"),
        ("코딩테스트 심화", "대기업 공채 기출 수준(삼성 SW역량테스트, 카카오/네이버 코테)까지 끌어올리기"),
        ("채용연계형 인턴 지원", "이 시기 인턴 경험이 하반기 공채에 결정적으로 유리하게 작용"),
    ]
)

content_slide(
    "STAGE 06", "4학년 2학기 — 본격 취업 시즌", "실전 지원 · 최종 정리",
    [
        ("자소서 · 이력서 · 포트폴리오 마무리", "프로젝트 경험을 숫자와 결과로 설명할 수 있게 정리"),
        ("기술면접 준비", "CS 전공 지식 + 프로젝트 경험을 논리적으로 설명하는 연습"),
        ("하반기 공채 · 상시채용 총력 지원", "인턴 경험이 있다면 전환형 채용부터 우선 확인"),
    ]
)

# =================================================================
# Slide 9: 인턴십 프로그램 총정리
# =================================================================
s = add_slide(WHITE)
header(s, "INTERNSHIP", "인턴십 · 채용연계 교육 프로그램", "3학년부터 방학마다 지원 대상으로 체크")

intern_rows = [
    [("삼성전자 DS/DX 대학생 인턴", "https://www.samsungcareers.com"), ("체험형 인턴", None), ("지원 3~4월 · 실습 7~8월", None),
     ("계열사 실무 체험, 직무적성검사(GSAT)+면접 절차", None)],
    [("SSAFY (삼성청년SW·AI아카데미)", "https://www.ssafy.com"), ("교육형(취업연계)", None), ("연 2회 모집(상/하반기)", None),
     ("1년 무료 SW교육, 협력사 채용 연계 다수", None)],
    [("네이버 부스트캠프", "https://boostcamp.connect.or.kr"), ("교육형", None), ("연 1~2회 기수 모집", None),
     ("웹/AI 트랙, 실무형 프로젝트 중심 무료 교육", None)],
    [("카카오테크 부트캠프", "https://kakaotechbootcamp.com"), ("교육형(취업연계)", None), ("연 1~2회 모집", None),
     ("국비지원, 카카오 계열 채용 연계 프로그램", None)],
    [("우아한테크코스", "https://techcourse.woowahan.com"), ("교육형", None), ("연 1회(하반기 모집)", None),
     ("5주 프리코스 + 10개월 정규과정, 배민 실무 학습", None)],
    [("SW마에스트로", "https://www.swmaestro.org"), ("정부지원 최상급 교육", None), ("연 1회(상반기 모집)", None),
     ("과기정통부 주관, 최정예 개발자 양성 · 우수 스타트업 연계", None)],
    [("42Seoul (이노베이션아카데미)", "https://42seoul.kr"), ("무료 교육(비학위)", None), ("상시 모집", None),
     ("학벌 무관, 동료학습 기반 프로젝트형 코딩 교육", None)],
]
add_table(
    s, 0.6, 1.85, 12.13, [2.7, 1.9, 2.15, 5.38],
    [0.55] + [0.66] * len(intern_rows), ["프로그램 (클릭 시 이동)", "유형", "시기(예시)", "특징"], intern_rows,
    body_size=11.5
)
add_text(s, 0.6, 6.95, 12.1, 0.4, "모집 시기는 매년 변동 — 지원 학기 초 공식 홈페이지에서 최신 일정 확인 필수", 11.5, MUTED, italic=True)

# =================================================================
# Slide 10: 인턴/채용 정보 사이트
# =================================================================
s = add_slide(WHITE)
header(s, "WEBSITES", "인턴 · 채용 정보 사이트 모음", "사이트명을 클릭하면 바로 이동")

site_rows = [
    [("사람인", "https://www.saramin.co.kr"), ("신입 · 인턴 채용 통합 검색, 국내 최대 채용 포털", None)],
    [("잡코리아", "https://www.jobkorea.co.kr"), ("신입 · 인턴 채용 통합 검색, 대기업·중견기업 다수", None)],
    [("원티드", "https://www.wanted.co.kr"), ("IT · 스타트업 특화, 예상 연봉 정보 공개 채용", None)],
    [("링커리어", "https://linkareer.com"), ("대학생 인턴 · 공모전 · 대외활동 특화 플랫폼", None)],
    [("프로그래머스 스쿨", "https://school.programmers.co.kr"), ("코딩테스트 연습 + 기업 채용 연계", None)],
    [("캠퍼스픽", "https://www.campuspick.com"), ("대학생 대외활동 · 인턴 정보 모음", None)],
    [("잡플래닛", "https://www.jobplanet.co.kr"), ("기업 리뷰 · 실제 연봉 정보 확인", None)],
    [("로켓펀치", "https://www.rocketpunch.com"), ("스타트업 채용 정보 특화", None)],
]
add_table(
    s, 0.6, 1.85, 12.13, [3.0, 9.13],
    [0.5] + [0.58] * len(site_rows), ["사이트", "용도"], site_rows,
    body_size=13
)

# =================================================================
# Field detail slides
# =================================================================
field_slide(
    "FIELD 01", "백엔드 개발자", "가장 견고한 채용 축 — 자격증보다 실무 역량이 핵심",
    [
        "정보처리기사, SQLD (필수는 아니지만 서류 가산점)",
        "Java/Spring 또는 Node.js/Python, RDB·NoSQL, REST API 설계, Git 협업",
        "코딩테스트 통과 실력, 배포까지 완료한 프로젝트, 기술면접 대응력",
        "네이버, 카카오, 쿠팡, 우아한형제들(배민), 토스, 라인, 당근마켓",
        "약 4,500만원 ~ 6,500만원 (스타트업~대기업, 기업별 편차 큼)",
        "학사면 충분 — 대학원 불필요, 실무 역량이 학력보다 중요",
    ]
)

field_slide(
    "FIELD 02", "AI · 데이터", "생성형 AI 붐으로 신입 공고 5년간 +162% 증가",
    [
        "ADsP(데이터분석 준전문가), 빅데이터분석기사, SQLD",
        "Python, PyTorch/TensorFlow, LLM · RAG 활용 경험, 통계 · 수학 기초",
        "캐글 · 논문 · 프로젝트 등 정량적 성과물, 생성형 AI 최신 트렌드 이해",
        "네이버, 카카오, 업스테이지, 뤼튼테크놀로지스, LG AI연구원, 삼성SDS",
        "ML엔지니어 약 4,500~6,500만원 / 데이터사이언티스트 약 3,300~4,500만원",
        "학사도 가능하나 R&D·연구 직무는 석사 이상 우대 비중이 높음",
    ]
)

field_slide(
    "FIELD 03", "정보보안", "장기 성장률 +28.5% 전망 — 보안 인력 부족 지속",
    [
        "정보보안기사 · 산업기사, CPPG, (경력 후) CISSP",
        "네트워크 · 시스템 보안, 모의해킹, Splunk 등 SIEM(관제) 툴 활용",
        "CTF 참가 경험, 보안 프로젝트 · 동아리 활동",
        "안랩, 이스트시큐리티, SK쉴더스, 금융권 보안팀, 공공 · 국방 보안기관",
        "약 4,000만원 ~ 5,000만원대",
        "학사면 충분 — 자격증 · 실무 경험이 학력보다 중요",
    ]
)

field_slide(
    "FIELD 04", "클라우드 · DevOps", "클라우드 전환 지속으로 안정적 수요",
    [
        "AWS SAA/DevOps, CKA(쿠버네티스), Azure · GCP 자격증",
        "Docker · Kubernetes, CI/CD 파이프라인, Terraform 등 IaC",
        "배포 자동화 구축 경험, 인프라 프로젝트 포트폴리오",
        "네이버클라우드, NHN클라우드, 카카오엔터프라이즈, 메가존클라우드, 베스핀글로벌",
        "약 4,000만원 ~ 5,500만원",
        "학사면 충분 — 자격증 · 실습 경험 위주로 평가",
    ]
)

# =================================================================
# 2026 유망 분야
# =================================================================
s = add_slide(NAVY)
add_text(s, 0.6, 0.5, 10.0, 0.6, "2026년 기준 유망 분야", 27, WHITE, bold=True, font="Cambria")
add_text(s, 0.6, 1.08, 10.0, 0.4, "채용 데이터가 가리키는 견고한 수요처", 13.5, ICE_TXT)

cards = [
    ("1위", "AI 엔지니어", "2026년 유망 직종 1순위, 전 산업으로 수요 확대"),
    ("+33.5%", "데이터 사이언티스트", "장기 성장률 전망"),
    ("+28.5%", "정보보안 분석가", "보안 인력 부족 지속"),
    ("수요 1위", "백엔드 · 풀스택", "AI 다음으로 가장 견고한 채용 축"),
]
cw, ch, cgap, cx0, cy0 = 2.75, 3.0, 0.35, 0.7, 1.95
for i, (big, label, sub) in enumerate(cards):
    x = cx0 + i * (cw + cgap)
    add_rect(s, x, cy0, cw, ch, NAVY_DARK, rounded=True)
    add_text(s, x + 0.15, cy0 + 0.3, cw - 0.3, 0.75, big, 25, AMBER, bold=True, font="Cambria", align=PP_ALIGN.CENTER)
    add_text(s, x + 0.15, cy0 + 1.18, cw - 0.3, 0.55, label, 15, WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x + 0.2, cy0 + 1.76, cw - 0.4, 1.1, sub, 11.5, ICE, align=PP_ALIGN.CENTER)

add_text(s, 0.7, 5.35, 11.5, 0.5,
         "핵심 역량 변화: 단순 코딩 능력 → AI 리터러시 · 데이터 이해력 · 문제 정의 능력",
         13, ICE_TXT, italic=True)
add_text(s, 0.7, 6.7, 11.5, 0.5,
         "2025~2026년 채용공고 절반 이상에 'LLM 경험', 'RAG 구축', 'GenAI 활용' 문구 포함",
         12, "8494C4", italic=True)

# =================================================================
# 대기업 초봉 스냅샷
# =================================================================
s = add_slide(WHITE)
header(s, "SALARY SNAPSHOT", "대기업 신입 초봉 스냅샷", "2025~2026 시장 조사 기준 — 성과급 포함 추정치, 참고용")

salary_rows = [
    [("SK하이닉스", None), ("약 1억원 이상 (성과급 포함)", None), ("반도체 특수 케이스, 업계 최고 수준", None)],
    [("삼성전자 DS/DX", None), ("약 6,000만원 ~ 9,000만원", None), ("성과급(OPI) 비중이 큼", None)],
    [("SK텔레콤", None), ("약 8,000만원대", None), ("통신 대기업 상위권", None)],
    [("네이버 · 카카오", None), ("약 5,000만원 ~ 6,000만원대", None), ("IT 플랫폼 대표 기업", None)],
    [("쿠팡 · 토스 등 유니콘", None), ("약 6,500만원 ~ 9,000만원", None), ("스톡옵션 포함 시 1억원 이상 사례도", None)],
    [("중견 · 스타트업 평균", None), ("약 3,500만원 ~ 4,500만원대", None), ("기업 규모 · 투자단계에 따라 편차 큼", None)],
]
add_table(
    s, 0.6, 1.95, 12.13, [3.3, 3.6, 5.23],
    [0.5] + [0.62] * len(salary_rows), ["기업 · 그룹", "신입 초봉 (약)", "비고"], salary_rows,
    body_size=12.5
)
add_text(s, 0.6, 6.85, 12.1, 0.45,
         "실제 금액은 기업 · 직무 · 성과급 정책에 따라 크게 달라짐 — 채용 시 잡플래닛 등에서 최신 정보 재확인 권장",
         11.5, MUTED, italic=True)

# =================================================================
# 종합 우선순위
# =================================================================
s = add_slide(WHITE)
add_text(s, 0.6, 0.45, 10.0, 0.6, "종합 우선순위", 28, NAVY, bold=True, font="Cambria")
add_text(s, 0.6, 1.05, 10.0, 0.4, "시간이 부족하면 이 순서대로", 14, MUTED)

items = [
    ("코딩테스트 실력", "가장 오래 걸리고 가장 확실한 무기. 2학년부터 꾸준히 누적"),
    ("포트폴리오 프로젝트 1~2개", "자격증보다 채용에 직접적 영향 — 배포까지 완료한 경험"),
    ("인턴십 경험", "3학년 방학부터 반복 지원. 붙기 전까지 광탈은 정상 과정"),
    ("자격증 (분야별 상세 11~14p)", "필수는 아니지만 있으면 서류에서 마이너스 없음"),
    ("CS 전공 지식", "코딩테스트 + 기술면접 모두에 직결 — 수업 자체가 준비 과정"),
]
top, bottom = 1.75, 7.0
row_h = (bottom - top) / len(items)
for i, (h, b) in enumerate(items):
    y = top + row_h * i
    fill = AMBER if i == 0 else NAVY
    add_badge(s, 0.6, y + (row_h - 0.55) / 2, 0.55, i + 1, fill, WHITE, size=15)
    add_text(s, 1.45, y + 0.03, 10.8, 0.4, h, 16.5, NAVY, bold=True)
    add_text(s, 1.45, y + 0.44, 11.0, row_h - 0.5, b, 12.5, MUTED)

OUT = r"C:\Users\SUNJIN\AppData\Local\Temp\claude\C--Users-SUNJIN-Desktop----\7621f797-e92a-470e-8842-3eb2bb406f40\scratchpad\roadmap_ppt\군복무_졸업_로드맵_상세.pptx"
prs.save(OUT)
print("saved", OUT)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
